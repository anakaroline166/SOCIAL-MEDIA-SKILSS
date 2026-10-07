import json, sys, copy
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from PIL import ImageFont
from lxml import etree

BASE, FONTS, OUTFILE = sys.argv[1:4]
data = json.load(open(f"{BASE}/texto.json")); texts, durs = data["texts"], data["durs"]
PX = 6350                                    # 1 px do slide (1920x1080) = 6350 EMU
prs = Presentation(); prs.slide_width, prs.slide_height = Emu(1920*PX), Emu(1080*PX)
blank = prs.slide_layouts[6]
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
for i, dur in enumerate(durs, 1):
    sl = prs.slides.add_slide(blank)
    sl.shapes.add_picture(f"{BASE}/slides/tela_{i:02d}.png", 0, 0, prs.slide_width, prs.slide_height)
    for t in [t for t in texts if t["slide"] == i]:
        path = f"{FONTS}/{'IntroBold' if t['bold'] else 'IntroBook'}.otf"
        pf = ImageFont.truetype(path, t["size"]); asc, desc = pf.getmetrics()
        w = pf.getlength(t["text"])
        tb = sl.shapes.add_textbox(Emu(int(t["x"]*PX)), Emu(int((t["y"]-asc)*PX)), Emu(int((w+30)*PX)), Emu(int((asc+desc)*PX)))
        tf = tb.text_frame; tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        r = tf.paragraphs[0].add_run(); r.text = t["text"]
        r.font.name = "Intro Bold" if t["bold"] else "Intro Book"
        r.font.size = Pt(t["size"]*0.5)
        r.font.color.rgb = RGBColor(*t["color"])
    # avanço automático + fade
    tr = etree.fromstring(f'<p:transition xmlns:p="{P}" spd="med" advTm="{dur*1000}"><p:fade/></p:transition>')
    clr = sl._element.find(f"{{{P}}}clrMapOvr")
    (clr.addnext(tr) if clr is not None else sl._element.append(tr))
# repetir em loop
for part in prs.part.package.iter_parts():
    if str(part.partname) == "/ppt/presProps.xml":
        part._blob = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
          f'<p:presentationPr xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="{P}">'
          '<p:showPr loop="1" showNarration="1"><p:present/><p:sldAll/><p:penClr><a:prstClr val="red"/></p:penClr></p:showPr></p:presentationPr>').encode()
prs.save(OUTFILE); print("pptx ok", len(durs), "slides")
